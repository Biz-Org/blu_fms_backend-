# -*- coding: utf-8 -*-

import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    Workbook = load_workbook = None


class MaintenanceRcaExcelImportWizard(models.TransientModel):
    _name = 'maintenance.rca.excel.import.wizard'
    _description = 'Import RCA & MOM master data from Excel'

    file = fields.Binary(string='Excel file', attachment=False)
    filename = fields.Char(string='Filename')
    request_id = fields.Many2one(
        'maintenance.request',
        string='Target work order',
        help='Optional: import only rows matching this work order reference.',
    )
    import_mode = fields.Selection(
        [
            ('create_or_update', 'Create or update work orders'),
            ('update_only', 'Update existing work orders only'),
        ],
        string='Import mode',
        default='create_or_update',
        required=True,
    )
    import_equipment = fields.Boolean(
        string='Import equipment master sheet',
        default=True,
        help='Create or update equipment from the Equipment Master sheet.',
    )
    replace_lines = fields.Boolean(
        string='Replace 5 Why & CAPA lines',
        default=True,
        help='When checked, existing 5 Why and corrective action lines are replaced by sheet data.',
    )
    import_log = fields.Text(string='Import log', readonly=True)

    # -------------------------------------------------------------------------
    # Template
    # -------------------------------------------------------------------------

    @api.model
    def _sheet_rows(self, worksheet):
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or '').strip() for h in rows[0]]
        data = []
        for row in rows[1:]:
            if not any(cell not in (None, '') for cell in row):
                continue
            item = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                if value is not None and not isinstance(value, str):
                    value = str(value).strip() if value != '' else ''
                else:
                    value = (value or '').strip()
                item[header.lower()] = value
            data.append(item)
        return data

    @api.model
    def _find_sheet(self, workbook, names):
        normalized = {name.strip().lower(): name for name in workbook.sheetnames}
        for candidate in names:
            key = candidate.strip().lower()
            if key in normalized:
                return workbook[normalized[key]]
        return None

    @api.model
    def _build_template_workbook(self):
        if not Workbook:
            raise UserError(_('openpyxl is required to generate Excel templates.'))

        wb = Workbook()
        guide = wb.active
        guide.title = 'Master Guide'
        guide['A1'] = 'Plant Maintenance — RCA & MOM Excel Master Data'
        guide['A1'].font = Font(bold=True, size=14)
        instructions = [
            '',
            'Fill the sheets below and upload from Analysis → Root Cause Analysis & MOM.',
            '',
            '1. Equipment Master — reference list of plant assets (optional create/update on import).',
            '2. RCA and MOM — work order header + root cause + meeting minutes.',
            '3. 5 Why Analysis — one row per Why step (same Work Order Ref).',
            '4. Corrective Actions — CAPA actions (Done = Yes / No).',
            '',
            'Work Order Ref must match an existing WO code (e.g. WO-224) or a new code to create.',
            'Fishbone format: Man: …; Machine: …; Method: …; Material: …; Environment: …',
        ]
        for idx, line in enumerate(instructions, start=2):
            guide[f'A{idx}'] = line

        header_fill = PatternFill('solid', fgColor='E8EEF7')
        bold = Font(bold=True)

        eq_ws = wb.create_sheet('Equipment Master')
        eq_headers = [
            'Equipment Name', 'Serial Number', 'Category', 'Location', 'Work Area',
        ]
        eq_ws.append(eq_headers)
        for col_idx, title in enumerate(eq_headers, start=1):
            cell = eq_ws.cell(row=1, column=col_idx)
            cell.font = bold
            cell.fill = header_fill
        equipment = self.env['maintenance.equipment'].sudo().search([], order='name', limit=100)
        for eq in equipment:
            eq_ws.append([
                eq.name or '',
                eq.serial_no or '',
                eq.category_id.name or '',
                eq.location or '',
                '',
            ])
        if not equipment:
            eq_ws.append([
                'Air Compressor 2', 'SN-DEMO-AC2', 'Compressors', 'Plant A', 'Utility Room',
            ])

        rca_ws = wb.create_sheet('RCA and MOM')
        rca_headers = [
            'Work Order Ref', 'Issue / Description', 'Equipment Serial', 'Maintenance Type',
            'Final Root Cause', 'Fishbone Summary', 'MOM Discussion', 'MOM Decision',
            'MOM Action Owner', 'Vibration', 'Bearing Temp', 'Failure Risk %',
            'Impact', 'Actual Hours', 'Work Area',
        ]
        rca_ws.append(rca_headers)
        for col_idx in range(1, len(rca_headers) + 1):
            cell = rca_ws.cell(row=1, column=col_idx)
            cell.font = bold
            cell.fill = header_fill
        rca_ws.append([
            'WO-224',
            'Breakdown — low pressure & high discharge temperature',
            'SN-DEMO-AC2',
            'corrective',
            'Failed pressure switch and worn coupling causing unstable pressure and overheating',
            'Man: missed weekly PM checklist; Machine: worn coupling and failed pressure switch; '
            'Method: no pre-start inspection; Material: degraded lubricant; Environment: high ambient dust',
            'Maintenance team reviewed compressor alarm history and MOM held with operations.',
            'Replace pressure switch, inspect coupling, schedule weekly PM checklist.',
            'Rahul — Machines Maintenance Team',
            'critical_high',
            'high',
            72,
            'high',
            3.5,
            'Utility Room',
        ])

        why_ws = wb.create_sheet('5 Why Analysis')
        why_headers = ['Work Order Ref', 'Step', 'Why Question', 'Because Answer']
        why_ws.append(why_headers)
        for col_idx in range(1, len(why_headers) + 1):
            cell = why_ws.cell(row=1, column=col_idx)
            cell.font = bold
            cell.fill = header_fill
        why_rows = [
            ('WO-224', 1, 'Why did discharge temperature rise?',
             'Aftercooler airflow was restricted'),
            ('WO-224', 2, 'Why was airflow restricted?',
             'Cooler fins were fouled with dust'),
            ('WO-224', 3, 'Why were fins fouled?',
             'Weekly cleaning was skipped on PM'),
            ('WO-224', 4, 'Why was cleaning skipped?',
             'Checklist not assigned to shift lead'),
            ('WO-224', 5, 'Why was checklist not assigned?',
             'No standard PM checklist for this asset'),
        ]
        for row in why_rows:
            why_ws.append(list(row))

        capa_ws = wb.create_sheet('Corrective Actions')
        capa_headers = ['Work Order Ref', 'Action', 'Done (Yes/No)']
        capa_ws.append(capa_headers)
        for col_idx in range(1, len(capa_headers) + 1):
            cell = capa_ws.cell(row=1, column=col_idx)
            cell.font = bold
            cell.fill = header_fill
        capa_ws.append([
            'WO-224', 'Implement weekly PM checklist for Air Compressor 2', 'Yes',
        ])
        capa_ws.append([
            'WO-224', 'Train operators on pre-start inspection', 'No',
        ])
        return wb

    @api.model
    def _workbook_to_attachment(self, wb, filename):
        buffer = io.BytesIO()
        wb.save(buffer)
        return self.env['ir.attachment'].sudo().create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(buffer.getvalue()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

    def action_download_template(self):
        self.ensure_one()
        wb = self._build_template_workbook()
        attachment = self._workbook_to_attachment(
            wb, 'Plant_Maintenance_RCA_Master_Template.xlsx'
        )
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    # -------------------------------------------------------------------------
    # Parsing helpers
    # -------------------------------------------------------------------------

    @api.model
    def _cell(self, row, *keys, default=''):
        for key in keys:
            val = row.get(key)
            if val not in (None, ''):
                return val
        return default

    @api.model
    def _parse_bool(self, value):
        text = str(value or '').strip().lower()
        return text in ('1', 'true', 'yes', 'y', 'done', 'complete', 'completed')

    @api.model
    def _parse_selection(self, value, mapping, default=None):
        text = str(value or '').strip().lower().replace(' ', '_').replace('-', '_')
        if text in mapping:
            return mapping[text]
        for key, result in mapping.items():
            if text == key.replace(' ', '_'):
                return result
        return default

    @api.model
    def _parse_float(self, value, default=0.0):
        if value in (None, ''):
            return default
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    @api.model
    def _parse_int(self, value, default=0):
        if value in (None, ''):
            return default
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default

    @api.model
    def _resolve_equipment(self, serial='', name=''):
        Equipment = self.env['maintenance.equipment'].sudo()
        serial = (serial or '').strip()
        name = (name or '').strip()
        if serial:
            eq = Equipment.search([('serial_no', '=', serial)], limit=1)
            if eq:
                return eq
        if name:
            eq = Equipment.search([('name', '=', name)], limit=1)
            if eq:
                return eq
        return Equipment.browse()

    @api.model
    def _find_request(self, work_order_ref):
        ref = (work_order_ref or '').strip()
        if not ref:
            return self.env['maintenance.request'].browse()
        Request = self.env['maintenance.request'].sudo()
        req = Request.search([('work_order_sequence', '=', ref)], limit=1)
        if req:
            return req
        return Request.search([('name', 'ilike', ref)], limit=1)

    @api.model
    def _default_stage_and_team(self):
        stage = self.env['maintenance.stage'].sudo().search([], limit=1)
        team = self.env['maintenance.team'].sudo().search([], limit=1)
        return stage, team

    @api.model
    def _import_equipment_rows(self, rows, log_lines):
        Equipment = self.env['maintenance.equipment'].sudo()
        Category = self.env['maintenance.equipment.category'].sudo()
        count = 0
        for row in rows:
            name = self._cell(row, 'equipment name', 'name')
            serial = self._cell(row, 'serial number', 'serial_no', 'serial')
            if not name and not serial:
                continue
            category_name = self._cell(row, 'category')
            location = self._cell(row, 'location')
            work_area = self._cell(row, 'work area', 'work_area')
            eq = self._resolve_equipment(serial=serial, name=name)
            category = Category.search([('name', '=', category_name)], limit=1) if category_name else False
            vals = {}
            if name:
                vals['name'] = name
            if serial:
                vals['serial_no'] = serial
            if category:
                vals['category_id'] = category.id
            if location:
                vals['location'] = location
            if eq:
                eq.write({k: v for k, v in vals.items() if v})
                log_lines.append(_('Equipment updated: %s (%s)') % (eq.name, eq.serial_no or '—'))
            else:
                if not name:
                    log_lines.append(_('Skipped equipment row — missing name for serial %s') % serial)
                    continue
                eq = Equipment.create(vals)
                log_lines.append(_('Equipment created: %s (%s)') % (eq.name, eq.serial_no or '—'))
            if work_area and eq:
                requests = self.env['maintenance.request'].sudo().search([
                    ('equipment_id', '=', eq.id),
                    ('work_area', 'in', (False, '')),
                ], limit=20)
                if requests:
                    requests.write({'work_area': work_area})
            count += 1
        return count

    @api.model
    def _import_rca_rows(self, rows, import_mode, target_request, log_lines):
        Request = self.env['maintenance.request'].sudo()
        stage, team = self._default_stage_and_team()
        count = 0
        maintenance_type_map = {
            'corrective': 'corrective',
            'cm': 'corrective',
            'preventive': 'preventive',
            'pm': 'preventive',
            'predictive': 'predictive',
            'pdm': 'predictive',
        }
        vibration_map = {
            'normal': 'normal',
            'critical_high': 'critical_high',
            'critical': 'critical_high',
            'high': 'critical_high',
        }
        bearing_map = {
            'normal': 'normal',
            'high': 'high',
        }
        impact_map = {
            'low': 'low',
            'medium': 'medium',
            'high': 'high',
        }

        for row in rows:
            wo_ref = self._cell(row, 'work order ref', 'work_order_ref', 'work order')
            if target_request and wo_ref and wo_ref != (target_request.work_order_sequence or ''):
                continue
            if target_request and not wo_ref:
                request = target_request
            else:
                request = self._find_request(wo_ref)

            if not request and import_mode == 'update_only':
                log_lines.append(_('Skipped (not found): %s') % (wo_ref or '—'))
                continue

            equipment_serial = self._cell(row, 'equipment serial', 'equipment_serial', 'serial number')
            equipment = self._resolve_equipment(serial=equipment_serial)
            issue = self._cell(row, 'issue / description', 'issue', 'description', 'name')
            vals = {
                'final_root_cause': self._cell(row, 'final root cause', 'final_root_cause'),
                'fishbone_summary': self._cell(row, 'fishbone summary', 'fishbone_summary'),
                'mom_discussion': self._cell(row, 'mom discussion', 'mom_discussion'),
                'mom_decision': self._cell(row, 'mom decision', 'mom_decision'),
                'mom_action_owner': self._cell(row, 'mom action owner', 'mom_action_owner'),
                'work_area': self._cell(row, 'work area', 'work_area'),
                'failure_risk_percent': self._parse_float(
                    self._cell(row, 'failure risk %', 'failure risk', 'failure_risk_percent')
                ),
                'actual_hours': self._parse_float(self._cell(row, 'actual hours', 'actual_hours')),
            }
            maint_type = self._parse_selection(
                self._cell(row, 'maintenance type', 'maintenance_type'),
                maintenance_type_map,
                default='corrective',
            )
            vals['maintenance_type'] = maint_type
            if maint_type == 'corrective':
                vals['failure_found_during_job'] = False
            else:
                vals['failure_found_during_job'] = True

            vibration = self._parse_selection(
                self._cell(row, 'vibration', 'vibration_level'), vibration_map, default='normal'
            )
            bearing = self._parse_selection(
                self._cell(row, 'bearing temp', 'bearing_temp', 'bearing_temp_state'),
                bearing_map,
                default='normal',
            )
            impact = self._parse_selection(
                self._cell(row, 'impact', 'impact_level'), impact_map, default='medium'
            )
            vals['vibration_level'] = vibration
            vals['bearing_temp_state'] = bearing
            vals['impact_level'] = impact

            if equipment:
                vals['equipment_id'] = equipment.id
            if issue:
                vals['name'] = issue

            if request:
                request.write({k: v for k, v in vals.items() if v not in (None, '', False) or k in (
                    'failure_risk_percent', 'actual_hours', 'failure_found_during_job',
                )})
                log_lines.append(_('RCA updated: %s') % (request.work_order_sequence or request.display_name))
            else:
                if not wo_ref:
                    log_lines.append(_('Skipped RCA row — missing Work Order Ref'))
                    continue
                create_vals = dict(vals)
                create_vals['work_order_sequence'] = wo_ref
                create_vals['name'] = issue or _('Imported maintenance issue')
                if stage:
                    create_vals['stage_id'] = stage.id
                if team:
                    create_vals['maintenance_team_id'] = team.id
                request = Request.create(create_vals)
                log_lines.append(_('RCA created: %s') % wo_ref)
            count += 1
        return count

    @api.model
    def _import_why_rows(self, rows, replace_lines, target_request, log_lines):
        grouped = {}
        for row in rows:
            wo_ref = self._cell(row, 'work order ref', 'work_order_ref')
            if target_request and wo_ref and wo_ref != (target_request.work_order_sequence or ''):
                continue
            if target_request and not wo_ref:
                request = target_request
            else:
                request = self._find_request(wo_ref)
            if not request:
                log_lines.append(_('5 Why skipped — work order not found: %s') % (wo_ref or '—'))
                continue
            grouped.setdefault(request.id, []).append(row)

        count = 0
        for request_id, request_rows in grouped.items():
            request = self.env['maintenance.request'].browse(request_id)
            if replace_lines:
                request.rca_line_ids.unlink()
            lines = []
            for row in sorted(
                request_rows,
                key=lambda r: self._parse_int(self._cell(r, 'step', 'sequence'), default=999),
            ):
                why = self._cell(row, 'why question', 'why', 'problem_statement')
                because = self._cell(row, 'because answer', 'because', 'cause_statement')
                if not why or not because:
                    continue
                seq = self._parse_int(self._cell(row, 'step', 'sequence'), default=len(lines) + 1)
                lines.append((0, 0, {
                    'sequence': seq or len(lines) + 1,
                    'problem_statement': why,
                    'cause_statement': because,
                }))
            if lines:
                request.write({'rca_line_ids': lines})
                count += len(lines)
                log_lines.append(
                    _('5 Why lines imported for %s: %s row(s)') % (
                        request.work_order_sequence or request.display_name, len(lines),
                    )
                )
        return count

    @api.model
    def _import_capa_rows(self, rows, replace_lines, target_request, log_lines):
        grouped = {}
        for row in rows:
            wo_ref = self._cell(row, 'work order ref', 'work_order_ref')
            if target_request and wo_ref and wo_ref != (target_request.work_order_sequence or ''):
                continue
            if target_request and not wo_ref:
                request = target_request
            else:
                request = self._find_request(wo_ref)
            if not request:
                log_lines.append(_('CAPA skipped — work order not found: %s') % (wo_ref or '—'))
                continue
            grouped.setdefault(request.id, []).append(row)

        count = 0
        for request_id, request_rows in grouped.items():
            request = self.env['maintenance.request'].browse(request_id)
            if replace_lines:
                request.corrective_action_ids.unlink()
            actions = []
            for row in request_rows:
                action_name = self._cell(row, 'action', 'name', 'corrective action')
                if not action_name:
                    continue
                actions.append((0, 0, {
                    'name': action_name,
                    'done': self._parse_bool(self._cell(row, 'done (yes/no)', 'done')),
                }))
            if actions:
                request.write({'corrective_action_ids': actions})
                count += len(actions)
                log_lines.append(
                    _('Corrective actions imported for %s: %s row(s)') % (
                        request.work_order_sequence or request.display_name, len(actions),
                    )
                )
        return count

    # -------------------------------------------------------------------------
    # Import action
    # -------------------------------------------------------------------------

    def action_import_excel(self):
        self.ensure_one()
        if not load_workbook:
            raise UserError(_('openpyxl is required to import Excel files.'))
        if not self.file:
            raise UserError(_('Please choose an Excel file (.xlsx) to import.'))

        try:
            workbook = load_workbook(filename=io.BytesIO(base64.b64decode(self.file)), data_only=True)
        except Exception as exc:
            raise UserError(_('Could not read the Excel file: %s') % exc) from exc

        log_lines = [_('Import started: %s') % (self.filename or 'uploaded file')]
        equipment_count = rca_count = why_count = capa_count = 0

        if self.import_equipment:
            eq_sheet = self._find_sheet(workbook, ['Equipment Master', 'Equipment'])
            if eq_sheet:
                equipment_count = self._import_equipment_rows(
                    self._sheet_rows(eq_sheet), log_lines
                )

        rca_sheet = self._find_sheet(workbook, ['RCA and MOM', 'RCA & MOM', 'RCA'])
        if rca_sheet:
            rca_count = self._import_rca_rows(
                self._sheet_rows(rca_sheet),
                self.import_mode,
                self.request_id,
                log_lines,
            )

        why_sheet = self._find_sheet(workbook, ['5 Why Analysis', '5 Why', 'Why Why'])
        if why_sheet:
            why_count = self._import_why_rows(
                self._sheet_rows(why_sheet),
                self.replace_lines,
                self.request_id,
                log_lines,
            )

        capa_sheet = self._find_sheet(workbook, ['Corrective Actions', 'CAPA', 'Corrective Action'])
        if capa_sheet:
            capa_count = self._import_capa_rows(
                self._sheet_rows(capa_sheet),
                self.replace_lines,
                self.request_id,
                log_lines,
            )

        touched = self.env['maintenance.request'].sudo().browse()
        if self.request_id:
            touched = self.request_id
        else:
            refs = set()
            for sheet_names in (
                ['RCA and MOM', 'RCA & MOM', 'RCA'],
                ['5 Why Analysis', '5 Why', 'Why Why'],
                ['Corrective Actions', 'CAPA', 'Corrective Action'],
            ):
                sheet = self._find_sheet(workbook, sheet_names)
                if not sheet:
                    continue
                for row in self._sheet_rows(sheet):
                    ref = self._cell(row, 'work order ref', 'work_order_ref')
                    if ref:
                        refs.add(ref)
            if refs:
                touched = self.env['maintenance.request'].sudo().search([
                    ('work_order_sequence', 'in', list(refs)),
                ])
        if touched:
            touched._recompute_rca_dashboard_html_fields()

        log_lines.append('')
        log_lines.append(_('Summary: %s equipment row(s), %s RCA row(s), %s Why line(s), %s CAPA row(s).') % (
            equipment_count, rca_count, why_count, capa_count,
        ))
        self.import_log = '\n'.join(log_lines)

        return {
            'type': 'ir.actions.act_window',
            'name': _('Import RCA Master Data'),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
