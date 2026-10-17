# -*- coding: utf-8 -*-

import base64
import io
import logging
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    Workbook = load_workbook = None


class MaintenanceEquipmentImportLog(models.Model):
    _name = 'maintenance.equipment.import.log'
    _description = 'Equipment Excel Upload'
    _order = 'upload_date desc, id desc'

    filename = fields.Char(string='Excel file', required=True, readonly=True)
    source_file = fields.Binary(string='Uploaded file', readonly=True, attachment=False)
    source_filename = fields.Char(string='Uploaded file name', readonly=True)
    user_id = fields.Many2one(
        'res.users',
        string='Uploaded by',
        default=lambda self: self.env.user,
        readonly=True,
    )
    upload_date = fields.Datetime(
        string='Upload date',
        default=fields.Datetime.now,
        readonly=True,
    )
    company_id = fields.Many2one(
        'res.company',
        string='Company',
        default=lambda self: self.env.company,
        readonly=True,
    )
    equipment_ids = fields.One2many(
        'maintenance.equipment',
        'excel_import_id',
        string='Equipment from this upload',
        readonly=True,
    )
    linked_equipment_ids = fields.Many2many(
        'maintenance.equipment',
        'maintenance_equipment_import_log_rel',
        'log_id',
        'equipment_id',
        string='Linked equipment',
        readonly=True,
    )
    equipment_count = fields.Integer(
        string='Equipment records',
        compute='_compute_equipment_count',
    )
    total_records = fields.Integer(string='Rows in file', readonly=True)
    imported_count = fields.Integer(string='Imported', readonly=True)
    skipped_count = fields.Integer(string='Skipped', readonly=True)
    failed_count = fields.Integer(string='Failed', readonly=True)
    state = fields.Selection(
        [
            ('success', 'Success'),
            ('partial', 'Partial'),
            ('failed', 'Failed'),
        ],
        string='Status',
        readonly=True,
    )
    summary_message = fields.Html(string='Summary', readonly=True, sanitize=False)
    log_message = fields.Text(string='Detailed log', readonly=True)
    error_report = fields.Binary(string='Error report', readonly=True, attachment=False)
    error_report_filename = fields.Char(string='Error report file', readonly=True)

    @api.depends('equipment_ids')
    def _compute_equipment_count(self):
        Equipment = self.env['maintenance.equipment'].sudo()
        for record in self:
            record.equipment_count = Equipment.search_count([
                ('excel_import_id', '=', record.id),
            ])

    def _equipment_from_upload(self):
        """All equipment from this upload (ignores record rules)."""
        self.ensure_one()
        Equipment = self.env['maintenance.equipment'].sudo()
        return Equipment.search([
            '|',
            ('excel_import_id', '=', self.id),
            ('id', 'in', self.linked_equipment_ids.ids),
        ])

    def _register_equipment(self, equipment):
        self.ensure_one()
        if equipment not in self.linked_equipment_ids:
            self.sudo().write({'linked_equipment_ids': [(4, equipment.id)]})

    def _remove_upload_equipment(self):
        """Remove or hide all equipment from this upload."""
        self.ensure_one()
        equipment = self._equipment_from_upload()
        deletable = equipment.filtered(lambda e: not e.maintenance_ids)
        protected = equipment - deletable
        deletable.unlink()
        if protected:
            protected.write({
                'excel_import_id': False,
                'is_excel_import': False,
                'active': False,
            })
        return len(equipment)

    def _check_remove_access(self):
        user = self.env.user
        if user.has_group('maintenance.group_equipment_manager'):
            return
        if user.has_group('maintenance.group_plant_maintenance_supervisor'):
            return
        for record in self:
            if record.user_id != user:
                raise UserError(_('You can only remove your own Excel uploads.'))

    def action_download_file(self):
        self.ensure_one()
        if not self.source_file:
            raise UserError(_('No file is stored for this upload.'))
        return {
            'type': 'ir.actions.act_url',
            'url': (
                f'/web/content/maintenance.equipment.import.log/{self.id}'
                f'/source_file/{self.source_filename or self.filename}?download=true'
            ),
            'target': 'self',
        }

    def action_download_error_report(self):
        self.ensure_one()
        if not self.error_report:
            return False
        return {
            'type': 'ir.actions.act_url',
            'url': (
                f'/web/content/maintenance.equipment.import.log/{self.id}'
                f'/error_report/{self.error_report_filename or "import_errors.xlsx"}?download=true'
            ),
            'target': 'self',
        }

    def action_remove_upload(self):
        """Delete this upload and remove its equipment from the Equipment screen."""
        self.ensure_one()
        self._check_remove_access()
        count = self._remove_upload_equipment()
        self.sudo().unlink()
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Upload removed'),
                'message': _('%s equipment record(s) removed from Equipment.') % count,
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }

    def unlink(self):
        self._check_remove_access()
        for record in self:
            record._remove_upload_equipment()
        companies = self.mapped('company_id').filtered(lambda c: c) or self.env.company
        res = super(MaintenanceEquipmentImportLog, self.sudo()).unlink()
        self.env['maintenance.analytics.dashboard']._invalidate_for_companies(companies)
        return res


class MaintenanceEquipmentImportWizard(models.TransientModel):
    _name = 'maintenance.equipment.import.wizard'
    _description = 'Import Equipment from Excel'

    MASTER_HEADERS = [
        'Equipment Name',
        'Equipment Code',
        'Serial Number',
        'Equipment Category',
        'Maintenance Team',
        'Owner',
        'Responsible Person',
        'Vendor',
        'Model',
        'Purchase Date',
        'Warranty Expiry Date',
        'Cost',
        'Location',
        'MTBF',
        'MTTR',
        'Status',
    ]

    file = fields.Binary(string='Excel file (.xlsx)', attachment=False)
    filename = fields.Char(string='Filename')
    result_message = fields.Html(string='Import result', readonly=True, sanitize=False)

    # -------------------------------------------------------------------------
    # Template
    # -------------------------------------------------------------------------

    @api.model
    def _header_styles(self):
        return {
            'bold': Font(bold=True),
            'fill': PatternFill('solid', fgColor='E8EEF7'),
        }

    @api.model
    def _style_header_row(self, worksheet, col_count):
        styles = self._header_styles()
        for col_idx in range(1, col_count + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = styles['bold']
            cell.fill = styles['fill']

    @api.model
    def _build_template_workbook(self):
        if not Workbook:
            raise UserError(_('openpyxl is required to generate Excel templates.'))

        wb = Workbook()
        guide = wb.active
        guide.title = 'Instructions'
        guide['A1'] = 'Equipment Excel Upload'
        guide['A1'].font = Font(bold=True, size=14)
        lines = [
            '',
            '1. Fill the "Equipment Master" sheet (Equipment Name is required).',
            '2. On Equipment screen click Import Excel and upload this file.',
            '3. All rows appear in Equipment.',
            '4. To remove them later: Equipment → Uploaded Excel → delete that upload.',
            '',
            'Dates: YYYY-MM-DD or DD/MM/YYYY',
            'Status: Active or Inactive',
        ]
        for idx, line in enumerate(lines, start=2):
            guide[f'A{idx}'] = line

        ws = wb.create_sheet('Equipment Master')
        ws.append(list(self.MASTER_HEADERS))
        self._style_header_row(ws, len(self.MASTER_HEADERS))

        Team = self.env['maintenance.team'].sudo()
        Category = self.env['maintenance.equipment.category'].sudo()
        User = self.env['res.users'].sudo()
        team = Team.search([], limit=1)
        category = Category.search([], limit=1)
        owner = User.search([('share', '=', False), ('active', '=', True)], limit=1)

        ws.append([
            'Air Compressor 2',
            'EQ-AC-002',
            'SN-IMPORT-001',
            category.name if category else 'Compressors',
            team.name if team else 'Machines Maintenance Team',
            owner.name if owner else 'Admin',
            owner.name if owner else 'Admin',
            'Atlas Copco',
            'GA37 VSD+',
            '2022-01-15',
            '2027-01-15',
            45000,
            'Plant A — Utility Room',
            180,
            2,
            'Active',
        ])
        return wb

    @api.model
    def _workbook_bytes(self, wb):
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def action_download_template(self):
        self.ensure_one()
        content = self._workbook_bytes(self._build_template_workbook())
        attachment = self.env['ir.attachment'].sudo().create({
            'name': 'Equipment_Upload_Template.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(content),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    # -------------------------------------------------------------------------
    # Parsing helpers
    # -------------------------------------------------------------------------

    @api.model
    def _sheet_rows(self, worksheet):
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or '').strip() for h in rows[0]]
        data = []
        for row_idx, row in enumerate(rows[1:], start=2):
            if not any(cell not in (None, '') for cell in row):
                continue
            item = {'_row_number': row_idx}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                if isinstance(value, datetime):
                    value = value.date().isoformat()
                elif value is not None and not isinstance(value, str):
                    value = str(value).strip() if value != '' else ''
                else:
                    value = (value or '').strip()
                item[header.strip().lower()] = value
            data.append(item)
        return data

    @api.model
    def _find_sheet(self, workbook, names):
        normalized = {name.strip().lower(): name for name in workbook.sheetnames}
        for candidate in names:
            key = candidate.strip().lower()
            if key in normalized:
                return workbook[normalized[key]]
        return workbook.active

    @api.model
    def _cell(self, row, *keys, default=''):
        for key in keys:
            val = row.get(key)
            if val not in (None, ''):
                return val
        return default

    @api.model
    def _parse_date(self, value):
        text = str(value or '').strip()
        if not text:
            return False
        if isinstance(value, datetime):
            return value.date()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        raise ValueError(_('Invalid date format: %s') % text)

    @api.model
    def _parse_float(self, value, default=0.0):
        if value in (None, ''):
            return default
        try:
            return float(str(value).replace(',', '').strip())
        except (TypeError, ValueError):
            raise ValueError(_('Invalid number: %s') % value)

    @api.model
    def _parse_int(self, value, default=0):
        if value in (None, ''):
            return default
        try:
            return int(float(str(value).replace(',', '').strip()))
        except (TypeError, ValueError):
            raise ValueError(_('Invalid integer: %s') % value)

    @api.model
    def _parse_status(self, value):
        text = str(value or 'active').strip().lower()
        if text in ('inactive', 'archived', 'archive', 'no', 'false', '0'):
            return False
        return True

    @api.model
    def _find_user(self, label):
        label = (label or '').strip()
        if not label:
            return self.env['res.users'].browse()
        User = self.env['res.users'].sudo()
        user = User.search([
            ('name', '=ilike', label),
            ('share', '=', False),
            ('active', '=', True),
        ], limit=1)
        if user:
            return user
        return User.search([
            ('login', '=ilike', label),
            ('share', '=', False),
            ('active', '=', True),
        ], limit=1)

    @api.model
    def _find_partner(self, label):
        label = (label or '').strip()
        if not label:
            return self.env['res.partner'].browse()
        return self.env['res.partner'].sudo().search([('name', '=ilike', label)], limit=1)

    @api.model
    def _find_category(self, label):
        label = (label or '').strip()
        if not label:
            return self.env['maintenance.equipment.category'].browse()
        Category = self.env['maintenance.equipment.category'].sudo()
        category = Category.search([('name', '=ilike', label)], limit=1)
        return category or Category.create({'name': label})

    @api.model
    def _find_team(self, label):
        label = (label or '').strip()
        if not label:
            return self.env['maintenance.team'].browse()
        return self.env['maintenance.team'].sudo().search([('name', '=ilike', label)], limit=1)

    @api.model
    def _find_existing_equipment(self, code, serial, name):
        Equipment = self.env['maintenance.equipment'].sudo()
        if code:
            match = Equipment.search([('equipment_code', '=', code)], limit=1)
            if match:
                return match
        if serial:
            match = Equipment.search([('serial_no', '=', serial)], limit=1)
            if match:
                return match
        if name:
            return Equipment.search([('name', '=ilike', name)], limit=1)
        return Equipment.browse()

    @api.model
    def _duplicate_exists(self, equipment_code, serial_no, exclude=None):
        Equipment = self.env['maintenance.equipment'].sudo()
        exclude_id = exclude.id if exclude else False
        if equipment_code:
            domain = [('equipment_code', '=', equipment_code)]
            if exclude_id:
                domain.append(('id', '!=', exclude_id))
            if Equipment.search_count(domain):
                return _('Duplicate equipment code: %s') % equipment_code
        if serial_no:
            domain = [('serial_no', '=', serial_no)]
            if exclude_id:
                domain.append(('id', '!=', exclude_id))
            if Equipment.search_count(domain):
                return _('Duplicate serial number: %s') % serial_no
        return False

    @api.model
    def _build_equipment_vals(self, row, import_log):
        name = self._cell(row, 'equipment name', 'name')
        if not name:
            raise ValueError(_('Equipment Name is mandatory'))

        equipment_code = self._cell(row, 'equipment code', 'equipment_code')
        serial_no = self._cell(row, 'serial number', 'serial_no', 'serial')
        category = self._find_category(self._cell(row, 'equipment category', 'category'))
        team = self._find_team(self._cell(row, 'maintenance team', 'maintenance_team'))
        owner = self._find_user(self._cell(row, 'owner', 'owner_user'))
        responsible = self._find_user(
            self._cell(row, 'responsible person', 'responsible', 'corrective responsible')
        )
        vendor = self._find_partner(self._cell(row, 'vendor', 'partner'))

        vals = {
            'name': name,
            'equipment_code': equipment_code or False,
            'serial_no': serial_no or False,
            'category_id': category.id if category else False,
            'maintenance_team_id': team.id if team else False,
            'owner_user_id': owner.id if owner else False,
            'corrective_responsible_user_id': responsible.id if responsible else (owner.id if owner else False),
            'preventive_responsible_user_id': responsible.id if responsible else False,
            'predictive_responsible_user_id': responsible.id if responsible else False,
            'partner_id': vendor.id if vendor else False,
            'model': self._cell(row, 'model') or False,
            'location': self._cell(row, 'location') or False,
            'active': self._parse_status(self._cell(row, 'status', default='Active')),
            'excel_import_id': import_log.id,
            'is_excel_import': True,
        }

        purchase_date = self._cell(row, 'purchase date', 'effective_date', 'purchase_date')
        if purchase_date:
            vals['effective_date'] = self._parse_date(purchase_date)

        warranty = self._cell(row, 'warranty expiry date', 'warranty_date', 'warranty date')
        if warranty:
            vals['warranty_date'] = self._parse_date(warranty)

        cost = self._cell(row, 'cost')
        if cost not in (None, ''):
            vals['cost'] = self._parse_float(cost)

        mtbf = self._cell(row, 'mtbf', 'expected mtbf')
        if mtbf not in (None, ''):
            vals['expected_mtbf'] = self._parse_int(mtbf)

        mttr = self._cell(row, 'mttr', 'expected mttr')
        if mttr not in (None, ''):
            vals['expected_mttr'] = self._parse_int(mttr)

        if team and not vals.get('owner_user_id') and team.member_ids:
            vals['owner_user_id'] = team.member_ids[0].id

        if 'effective_date' not in vals:
            vals['effective_date'] = fields.Date.context_today(self)

        return vals

    @api.model
    def _build_error_report(self, error_rows):
        if not Workbook:
            return b'', False
        wb = Workbook()
        ws = wb.active
        ws.title = 'Import Errors'
        headers = ['Row', 'Equipment Name', 'Equipment Code', 'Serial Number', 'Error Reason']
        ws.append(headers)
        self._style_header_row(ws, len(headers))
        for item in error_rows:
            ws.append([
                item.get('row'),
                item.get('name', ''),
                item.get('code', ''),
                item.get('serial', ''),
                item.get('error', ''),
            ])
        return self._workbook_bytes(wb), 'Equipment_Import_Errors.xlsx'

    # -------------------------------------------------------------------------
    # Import
    # -------------------------------------------------------------------------

    def action_import_excel(self):
        self.ensure_one()
        if not load_workbook:
            raise UserError(_('openpyxl is required to import Excel files.'))
        if not self.file:
            raise UserError(_('Please choose an Excel file (.xlsx) to upload.'))

        try:
            workbook = load_workbook(
                filename=io.BytesIO(base64.b64decode(self.file)),
                data_only=True,
            )
        except Exception as exc:
            raise UserError(_('Could not read the Excel file: %s') % exc) from exc

        sheet = self._find_sheet(workbook, ['Equipment Master', 'Equipment', 'Master Data'])
        rows = self._sheet_rows(sheet)

        import_log = self.env['maintenance.equipment.import.log'].sudo().create({
            'filename': self.filename or 'uploaded_file.xlsx',
            'source_file': self.file,
            'source_filename': self.filename or 'uploaded_file.xlsx',
            'total_records': len(rows),
        })

        Equipment = self.env['maintenance.equipment'].sudo()
        imported = skipped = failed = 0
        log_lines = []
        error_rows = []

        for row in rows:
            row_no = row.get('_row_number', '?')
            name = self._cell(row, 'equipment name', 'name')
            code = self._cell(row, 'equipment code', 'equipment_code')
            serial = self._cell(row, 'serial number', 'serial_no', 'serial')
            try:
                existing = self._find_existing_equipment(code, serial, name)
                vals = self._build_equipment_vals(row, import_log)

                if existing:
                    if existing.excel_import_id and existing.excel_import_id != import_log:
                        skipped += 1
                        log_lines.append(
                            _('Row %s skipped: already linked to upload "%s"')
                            % (row_no, existing.excel_import_id.filename)
                        )
                        continue
                    if existing.maintenance_ids:
                        skipped += 1
                        log_lines.append(
                            _('Row %s skipped: "%s" has work orders')
                            % (row_no, existing.name)
                        )
                        continue
                    if not existing.active:
                        vals['active'] = True
                    existing.write(vals)
                    import_log._register_equipment(existing)
                    imported += 1
                    log_lines.append(
                        _('Row %s linked: %s') % (row_no, existing.name)
                    )
                    continue

                dup_reason = self._duplicate_exists(code, serial)
                if dup_reason:
                    skipped += 1
                    log_lines.append(_('Row %s skipped: %s') % (row_no, dup_reason))
                    continue

                equipment = Equipment.create(vals)
                import_log._register_equipment(equipment)
                imported += 1
                log_lines.append(
                    _('Row %s imported: %s') % (row_no, equipment.name)
                )
            except Exception as exc:
                failed += 1
                reason = str(exc)
                log_lines.append(_('Row %s failed: %s') % (row_no, reason))
                error_rows.append({
                    'row': row_no,
                    'name': name,
                    'code': code,
                    'serial': serial,
                    'error': reason,
                })
                _logger.warning('Equipment import row %s failed: %s', row_no, reason)

        if imported and not failed:
            state = 'success'
        elif imported:
            state = 'partial'
        else:
            state = 'failed'

        summary_html = (
            '<div class="alert alert-success" role="status">'
            f'<strong>{_("Upload complete")}</strong><br/>'
            f'{_("Equipment from Excel")}: {imported}<br/>'
            f'{_("Skipped")}: {skipped}<br/>'
            f'{_("Failed")}: {failed}<br/>'
            f'<em>{_("Delete this upload under Uploaded Excel to remove these equipment records.")}</em>'
            '</div>'
        )
        if not imported:
            summary_html += (
                '<div class="alert alert-warning mt-2" role="status">'
                f'{_("No rows were imported. Check Equipment Name column and use Download template.")}'
                '</div>'
            )
        elif skipped:
            summary_html += (
                '<div class="alert alert-info mt-2" role="status">'
                f'{_("Some rows were skipped because they duplicate equipment from another upload.")}'
                '</div>'
            )
        if failed:
            summary_html += (
                '<div class="alert alert-warning mt-2" role="status">'
                f'{_("Some rows failed. Download the error report for details.")}'
                '</div>'
            )

        error_report = False
        error_filename = False
        if error_rows:
            error_report, error_filename = self._build_error_report(error_rows)

        import_log.write({
            'imported_count': imported,
            'skipped_count': skipped,
            'failed_count': failed,
            'state': state,
            'summary_message': summary_html,
            'log_message': '\n'.join(log_lines),
            'error_report': base64.b64encode(error_report) if error_report else False,
            'error_report_filename': error_filename,
        })

        if imported:
            self.env['maintenance.analytics.dashboard']._invalidate_for_companies(self.env.company)

        self.write({'result_message': summary_html})

        if imported:
            return {
                'type': 'ir.actions.act_window',
                'name': _('Equipment from Excel'),
                'res_model': 'maintenance.equipment',
                'view_mode': 'kanban,list,form',
                'domain': [('excel_import_id', '=', import_log.id)],
                'target': 'current',
            }

        return {
            'type': 'ir.actions.act_window',
            'name': _('Import Excel'),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_open_uploads(self):
        return self.env.ref('maintenance.action_maintenance_equipment_import_log').read()[0]
